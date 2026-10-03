"""Оқушы толтыратын зерттеу кестесі: Zertteu_kestesi.xlsx
python3 make_xlsx.py шығыс.xlsx [толтырылған_ескі.xlsx]  – ескі кестедегі деректер жаңасына көшіріледі."""
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from sandar import CORPUS, TALE_TYPES, CATS, FUNCS, NUMBERS, SACRED, N_ROWS

OUT = sys.argv[1] if len(sys.argv) > 1 else "Zertteu_kestesi.xlsx"
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DCE6F1")
FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)


def header(ws, row, titles, widths):
    for j, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row, j, t)
        c.fill, c.font, c.alignment, c.border = HEAD, BOLD, CEN, BOX
        if w:
            ws.column_dimensions[c.column_letter].width = w


def put(ws, r, j, v=None, fill=False, center=False):
    c = ws.cell(r, j, v)
    c.border = BOX
    c.alignment = CEN if center else WRAP
    if fill:
        c.fill = FILL
    return c


wb = Workbook()
ws = wb.active
ws.title = "Нұсқау"
ws.column_dimensions["A"].width = 115
lines = [
    ("ЗЕРТТЕУ КЕСТЕСІН ТОЛТЫРУ НҰСҚАУЫ", True),
    ("Сары ұяшықтарды ғана толтырыңыз. «Қорытынды» парағы өздігінен есептеледі.", False),
    ("", False),
    ("1. «Ертегілер» парағындағы 14 ертегіні оқыңыз. Қай кітаптан (баспасы, жылы) оқығаныңызды жазып, «Оқылды» бағанына «иә» деп белгілеңіз.", False),
    ("2. Ертегіде үш ағайынды, үш сынақ, үш рет қайталанатын оқиға болса, «Үштік құрылым» бағанына «иә» деп, мысалын жазыңыз.", False),
    ("3. Ертегіде кездескен ӘР САНДЫ «Сандар» парағына жазыңыз: бір сан – бір жол.", False),
    ("   • «Сан» бағанына санды цифрмен жазыңыз (жеті → 7, қырық → 40, мың → 1000).", False),
    ("   • «Мәтіндегі тіркес» бағанына санмен бірге келген сөздерді жазыңыз: «жеті басты дәу», «үш ұлы бар екен».", False),
    ("   • «Нені білдіреді» бағанында тізімнен таңдаңыз: адам, уақыт, жануар, зат, жер-қашықтық, дене мүшесі, басқа.", False),
    ("4. Санның қызметін белгілеңіз:", False),
    *[(f"     {a} – {b}: {c}", False) for a, b, c in FUNCS],
    ("5. «Бірінші», «екінші» сияқты реттік сандарды да жазыңыз (бірінші → 1). «Бір күні», «бір шал» сияқты «бір» сөзін жазбаңыз – ол сан емес, «белгісіз біреу» деген мағынада.", False),
    ("6. Күмәнді жағдайды ғылыми жетекшіңізбен ақылдасып шешіңіз және шешімді зерттеу күнделігіне жазыңыз.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.alignment = WRAP
    c.font = Font(bold=b, size=12 if b else 11)

# ---- Ертегілер ----
wc = wb.create_sheet("Ертегілер")
header(wc, 1, ["№", "Ертегі", "Түрі", "Оқыған басылым (кітап, баспа, жылы)", "Оқылды (иә/жоқ)", "Үштік құрылым (иә/жоқ)", "Үштік құрылымның мысалы"],
       [5, 30, 15, 42, 11, 11, 45])
dv_yes = DataValidation(type="list", formula1='"иә,жоқ"', allow_blank=True)
wc.add_data_validation(dv_yes)
for i, (t, g) in enumerate(CORPUS, 1):
    r = i + 1
    put(wc, r, 1, i, center=True); put(wc, r, 2, t); put(wc, r, 3, g, center=True)
    put(wc, r, 4, fill=True)
    dv_yes.add(put(wc, r, 5, fill=True, center=True))
    dv_yes.add(put(wc, r, 6, fill=True, center=True))
    put(wc, r, 7, fill=True)
CORP_LAST = len(CORPUS) + 1

# ---- Анықтамалық ----
wl = wb.create_sheet("Тізімдер")
header(wl, 1, ["Нені білдіреді", "Қызметі"], [24, 12])
for i, c in enumerate(CATS, 2):
    wl.cell(i, 1, c)
for i, f in enumerate(FUNCS, 2):
    wl.cell(i, 2, f[0])

# ---- Сандар ----
wn = wb.create_sheet("Сандар")
header(wn, 1, ["№", "Ертегі", "Түрі", "Сан", "Мәтіндегі тіркес", "Нені білдіреді", "Қызметі (Н/Ф)", "Беті"], [5, 30, 14, 8, 44, 18, 10, 7])
wn.freeze_panes = "B2"
dv_tale = DataValidation(type="list", formula1=f"=Ертегілер!$B$2:$B${CORP_LAST}", allow_blank=True)
dv_cat = DataValidation(type="list", formula1=f"=Тізімдер!$A$2:$A${len(CATS) + 1}", allow_blank=True)
dv_func = DataValidation(type="list", formula1='"Н,Ф"', allow_blank=True)
dv_num = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
for dv in (dv_tale, dv_cat, dv_func, dv_num):
    wn.add_data_validation(dv)
for i in range(1, N_ROWS + 1):
    r = i + 1
    put(wn, r, 1, i, center=True)
    dv_tale.add(put(wn, r, 2, fill=True))
    put(wn, r, 3, f'=IFERROR(VLOOKUP(B{r},Ертегілер!$B$2:$C${CORP_LAST},2,0),"")', center=True)
    dv_num.add(put(wn, r, 4, fill=True, center=True))
    put(wn, r, 5, fill=True)
    dv_cat.add(put(wn, r, 6, fill=True))
    dv_func.add(put(wn, r, 7, fill=True, center=True))
    put(wn, r, 8, fill=True, center=True)
NL = N_ROWS + 1
RNG = lambda col: f"Сандар!${col}$2:${col}${NL}"

# ---- Қорытынды ----
wq = wb.create_sheet("Қорытынды")
for col, w in zip("ABCDE", [30, 14, 14, 14, 14]):
    wq.column_dimensions[col].width = w
row = 1
wq.cell(row, 1, "Сандардың кездесуі").font = BOLD
row += 1
header(wq, row, ["Сан", "Барлығы", "Нақты (Н)", "Формула (Ф)"], [None] * 4)
first = row
for n in NUMBERS:
    row += 1
    put(wq, row, 1, n, center=True)
    put(wq, row, 2, f"=COUNTIF({RNG('D')},A{row})", center=True)
    put(wq, row, 3, f"=COUNTIFS({RNG('D')},A{row},{RNG('G')},\"Н\")", center=True)
    put(wq, row, 4, f"=COUNTIFS({RNG('D')},A{row},{RNG('G')},\"Ф\")", center=True)
last = row
row += 1
put(wq, row, 1, "Басқа сандар", center=True)
put(wq, row, 2, f"=B{row + 1}-SUM(B{first + 1}:B{last})", center=True)
put(wq, row, 3, f"=C{row + 1}-SUM(C{first + 1}:C{last})", center=True)
put(wq, row, 4, f"=D{row + 1}-SUM(D{first + 1}:D{last})", center=True)
row += 1
put(wq, row, 1, "Барлығы").font = BOLD
put(wq, row, 2, f"=COUNT({RNG('D')})", center=True)
put(wq, row, 3, f"=COUNTIF({RNG('G')},\"Н\")", center=True)
put(wq, row, 4, f"=COUNTIF({RNG('G')},\"Ф\")", center=True)
total_row = row
ch = BarChart(); ch.type = "col"; ch.title = "Ертегілердегі сандар"; ch.grouping = "stacked"; ch.overlap = 100
ch.add_data(Reference(wq, min_col=3, max_col=4, min_row=first, max_row=last), titles_from_data=True)
ch.set_categories(Reference(wq, min_col=1, min_row=first + 1, max_row=last)); ch.height = 8; ch.width = 18
wq.add_chart(ch, "G1")

row += 2
wq.cell(row, 1, "«Киелі» сандар (" + ", ".join(map(str, SACRED)) + ")").font = BOLD
put(wq, row + 1, 1, "Саны")
put(wq, row + 1, 2, "=" + "+".join(f"COUNTIF({RNG('D')},{n})" for n in SACRED), center=True)
put(wq, row + 2, 1, "Барлық сандағы үлесі, %")
put(wq, row + 2, 2, f'=IFERROR(ROUND(B{row + 1}/B{total_row}*100,1),"")', center=True)
row += 4

wq.cell(row, 1, "Сан нені білдіреді").font = BOLD
row += 1
header(wq, row, ["Нені білдіреді", "Саны"], [None] * 2)
c_first = row
for c in CATS:
    row += 1
    put(wq, row, 1, c)
    put(wq, row, 2, f"=COUNTIF({RNG('F')},A{row})", center=True)
ch = BarChart(); ch.type = "bar"; ch.title = "Сан нені білдіреді"
ch.add_data(Reference(wq, min_col=2, min_row=c_first, max_row=row), titles_from_data=True)
ch.set_categories(Reference(wq, min_col=1, min_row=c_first + 1, max_row=row)); ch.height = 7; ch.width = 14; ch.legend = None
wq.add_chart(ch, "G18")

row += 2
wq.cell(row, 1, "Ертегі түрлері бойынша").font = BOLD
row += 1
header(wq, row, ["Түрі", "Сандар саны", "Оның ішінде Ф", "Үштік құрылым (ертегі саны)"], [None] * 4)
for t in TALE_TYPES:
    row += 1
    put(wq, row, 1, t)
    put(wq, row, 2, f"=COUNTIF({RNG('C')},A{row})", center=True)
    put(wq, row, 3, f"=COUNTIFS({RNG('C')},A{row},{RNG('G')},\"Ф\")", center=True)
    put(wq, row, 4, f"=COUNTIFS(Ертегілер!$C$2:$C${CORP_LAST},A{row},Ертегілер!$F$2:$F${CORP_LAST},\"иә\")", center=True)

if len(sys.argv) > 2:
    src = load_workbook(sys.argv[2], data_only=True)
    for name, rows, cols in [("Ертегілер", range(2, CORP_LAST + 1), (4, 5, 6, 7)), ("Сандар", range(2, NL + 1), (2, 4, 5, 6, 7, 8))]:
        for r in rows:
            for c in cols:
                v = src[name].cell(r, c).value
                wb[name].cell(r, c).value = v.strip() if isinstance(v, str) else v

for sh in wb.worksheets:
    sh.page_setup.orientation = "landscape"
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
wb.save(OUT)
print("written", OUT)
