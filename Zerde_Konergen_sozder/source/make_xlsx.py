"""Оқушы толтыратын зерттеу кестесі: Zertteu_kestesi.xlsx
python3 make_xlsx.py шығыс.xlsx [толтырылған_ескі.xlsx]  – ескі кестедегі деректер жаңасына көшіріледі."""
import json, os, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from zertteu import (GROUPS, KINDS, WORDS, DICT_MARKS, CONTEXTS, STATUS, RARE, SOURCE,
                     TEST_WORDS, GENS, N_PER_GEN, KNOW)

OUT = sys.argv[1] if len(sys.argv) > 1 else "Zertteu_kestesi.xlsx"
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DCE6F1")
FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)
GNAME = dict(GROUPS)
KNAME = {k: n for k, n, _ in KINDS}


def header(ws, row, titles, widths):
    for j, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row, j, t)
        c.fill, c.font, c.alignment, c.border = HEAD, BOLD, CEN, BOX
        if w:
            ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 45


def put(ws, r, j, v=None, fill=False, center=False):
    c = ws.cell(r, j, v)
    c.border = BOX
    c.alignment = CEN if center else WRAP
    if fill:
        c.fill = FILL
    return c


def dv_list(ws, items):
    dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Қате мән", error="Тек мына мәндердің бірін жазыңыз: " + ", ".join(items))
    ws.add_data_validation(dv)
    return dv


wb = Workbook()
ws = wb.active
ws.title = "Нұсқау"
ws.column_dimensions["A"].width = 120
lines = [
    ("ЗЕРТТЕУ КЕСТЕСІН ТОЛТЫРУ НҰСҚАУЫ", True),
    ("Сары ұяшықтарды ғана толтырыңыз. «Мәртебе» бағаны мен «Қорытынды» парағы өздігінен есептеледі. Парақтарды, бағандарды өшірмеңіз.", False),
    ("", False),
    ("1-КЕЗЕҢ. Сөздікпен тексеру («Сөздер» парағы)", True),
    ("1. Әр сөзді «Қазақ әдеби тілінің сөздігінен» (15 томдық) тауып, оның жанындағы белгіні жазыңыз: «көн.» (көнерген), «тар.» (тарихи), белгі жоқ немесе сөздікте жоқ.", False),
    ("2. Томын көрсетіңіз. Сөздікте берілген анықтаманы «Мағынасы» бағанына көшіріп алуға болады.", False),
    ("", False),
    ("2-КЕЗЕҢ. БАҚ мониторингі («БАҚ мониторингі» парағы)", True),
    ("3. Google іздеу жолына: site:egemen.kz \"сауыт\" деп жазыңыз. «Құралдар» → «Кез келген уақыт» → «Соңғы жыл» таңдаңыз. Шыққан нәтиже санын жазыңыз (мысалы, «Шамамен 45 нәтиже» болса – 45). Нәтиже жоқ болса – 0.", False),
    ("4. Барлық сөзді бір-екі күн ішінде тексеріңіз, күнін жазыңыз.", False),
    ("5. Ең жаңа мақаланы ашып, сөз ТҰРҒАН сөйлемді көшіріп алыңыз (мақаланың бірінші сөйлемін емес) және сілтемесін жазыңыз. Сөйлемнің мәнмәтінін бір кодпен белгілеңіз:", False),
    *[(f"     {a} – {b}: {c}", False) for a, b, c in CONTEXTS],
    ("", False),
    ("3-КЕЗЕҢ. Үш буынға сөз тану тесті («Тест» парағы)", True),
    (f"6. Әр буыннан {N_PER_GEN} адам: оқушылар, ата-аналар, ата-әжелер. Келісімін алыңыз (оқушылар үшін – ата-анасының келісімі). Есімдерін жазбаңыз – тек код.", False),
    ("7. Әр сөзді атап: «Бұл сөздің мағынасы қандай?» деп сұраңыз. 2 – дұрыс түсіндірді, 1 – шамамен түсіндірді, 0 – білмейді. Сөздің мағынасын айтып, көмектеспеңіз.", False),
    ("8. Сосын: «Бұл сөзді соңғы бір жылда естідіңіз бе немесе өзіңіз қолдандыңыз ба?» деп сұраңыз: 1 – иә, 0 – жоқ.", False),
    ("", False),
    ("МӘРТЕБЕ ЕРЕЖЕСІ (жетекшіммен келісілген)", True),
    (f"Ұ – нәтиже саны {RARE}-тен аз болса; Ж – әйтпесе, мысал «Ж» болса; А – мысал «А» болса; Т – қалған жағдайда («Т» не «Б»).", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.alignment = WRAP
    c.font = Font(bold=b, size=12 if b else 11)

# ---- Сөздер ----
wsz = wb.create_sheet("Сөздер")
header(wsz, 1, ["№", "Сөз", "Тақырыптық топ", "Түрі (алдын ала)", "Мағынасы", "Сөздіктегі белгі", "Томы"], [5, 14, 26, 14, 52, 14, 12])
wsz.freeze_panes = "C2"
dv_mark = dv_list(wsz, DICT_MARKS)
for i, (w, g, k, m) in enumerate(WORDS, 1):
    r = i + 1
    for j, v in enumerate([i, w, GNAME[g], KNAME[k], m], 1):
        put(wsz, r, j, v, center=j in (1, 4))
    dv_mark.add(put(wsz, r, 6, fill=True, center=True))
    put(wsz, r, 7, fill=True, center=True)
W_LAST = len(WORDS) + 1

# ---- БАҚ мониторингі ----
wm = wb.create_sheet("БАҚ мониторингі")
s1 = SOURCE[0]
header(wm, 1, ["№", "Сөз", f"{s1}: нәтиже саны", f"{s1}: мәнмәтін (код)", "Ең жаңа мысал (сөйлем)", "Сілтеме", "Тексерген күні", "Мәртебе"],
       [5, 13, 12, 12, 55, 30, 12, 10])
wm.freeze_panes = "C2"
dv_ctx = dv_list(wm, [c for c, _, _ in CONTEXTS])
for i, (w, _, _, _) in enumerate(WORDS, 1):
    r = i + 1
    put(wm, r, 1, i, center=True); put(wm, r, 2, w)
    put(wm, r, 3, fill=True, center=True); dv_ctx.add(put(wm, r, 4, fill=True, center=True))
    put(wm, r, 5, fill=True); put(wm, r, 6, fill=True); put(wm, r, 7, fill=True, center=True)
    put(wm, r, 8, f'=IF(C{r}="","",IF(N(C{r})<{RARE},"Ұ",IF(D{r}="Ж","Ж",IF(D{r}="А","А","Т"))))', center=True).font = BOLD

# ---- Тест ----
wt = wb.create_sheet("Тест")
titles = ["Код", "Буын"]
for w in TEST_WORDS:
    titles += [f"{w}: білу (0/1/2)", f"{w}: естиді (1/0)"]
header(wt, 1, titles, [7, 22] + [10] * (2 * len(TEST_WORDS)))
wt.row_dimensions[1].height = 60
wt.freeze_panes = "C2"
dv_k = dv_list(wt, ["0", "1", "2"]); dv_u = dv_list(wt, ["0", "1"])
r = 1
for code, gname in GENS:
    for i in range(1, N_PER_GEN + 1):
        r += 1
        put(wt, r, 1, f"{code}{i:02d}", center=True); put(wt, r, 2, gname)
        for k in range(len(TEST_WORDS)):
            dv_k.add(put(wt, r, 3 + 2 * k, fill=True, center=True))
            dv_u.add(put(wt, r, 4 + 2 * k, fill=True, center=True))
T_LAST = r

# ---- Қорытынды ----
wq = wb.create_sheet("Қорытынды")
for col, w in zip("ABCDEFG", [30, 12, 12, 12, 12, 12, 12]):
    wq.column_dimensions[col].width = w
row = 1
wq.cell(row, 1, "Сөздердің мәртебесі тақырыптық топтар бойынша").font = BOLD
row += 1
header(wq, row, ["Топ"] + [n for _, n, _ in STATUS] + ["Барлығы"], [None] * 6)
st_first = row
for gi, (g, gname) in enumerate(GROUPS):
    row += 1
    put(wq, row, 1, gname)
    idx = [i + 2 for i, w in enumerate(WORDS) if w[1] == g]
    for j, (code, _, _) in enumerate(STATUS, 2):
        put(wq, row, j, "=" + "+".join(f'(\'БАҚ мониторингі\'!H{x}="{code}")' for x in idx), center=True)
    put(wq, row, 6, f"=SUM(B{row}:E{row})", center=True)
row += 1
put(wq, row, 1, "Барлығы").font = BOLD
for j in range(2, 7):
    c = get_column_letter(j)
    put(wq, row, j, f"=SUM({c}{st_first + 1}:{c}{row - 1})", center=True)
ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100; ch.title = "Сөздердің бүгінгі мәртебесі"
ch.add_data(Reference(wq, min_col=2, max_col=5, min_row=st_first, max_row=row - 1), titles_from_data=True)
ch.set_categories(Reference(wq, min_col=1, min_row=st_first + 1, max_row=row - 1)); ch.height = 8; ch.width = 18
wq.add_chart(ch, "I1")

row += 2
wq.cell(row, 1, "Үш буын тесті: мағынасын дұрыс түсіндіргендер үлесі, %").font = BOLD
row += 1
header(wq, row, ["Сөз"] + [n for _, n in GENS], [None] * 4)
t_first = row
for k, w in enumerate(TEST_WORDS):
    row += 1
    put(wq, row, 1, w)
    col = get_column_letter(3 + 2 * k)
    for j, (code, _) in enumerate(GENS, 2):
        rng_c, rng_v = f"Тест!$A$2:$A${T_LAST}", f"Тест!${col}$2:${col}${T_LAST}"
        put(wq, row, j, f'=IFERROR(ROUND(COUNTIFS({rng_c},"{code}*",{rng_v},2)/COUNTIFS({rng_c},"{code}*",{rng_v},"<>")*100,1),"")', center=True)
ch = BarChart(); ch.type = "col"; ch.title = "Сөзді дұрыс түсіндіргендер, %"
ch.add_data(Reference(wq, min_col=2, max_col=4, min_row=t_first, max_row=row), titles_from_data=True)
ch.set_categories(Reference(wq, min_col=1, min_row=t_first + 1, max_row=row)); ch.height = 8; ch.width = 18
wq.add_chart(ch, "I18")

# ---- Ескі кестеден көшіру ----
if len(sys.argv) > 2:
    src = load_workbook(sys.argv[2], data_only=True)
    for name, rows, cols in [("Сөздер", range(2, W_LAST + 1), (5, 6, 7)),
                             ("БАҚ мониторингі", range(2, W_LAST + 1), range(3, 8)),
                             ("Тест", range(2, T_LAST + 1), range(3, 3 + 2 * len(TEST_WORDS)))]:
        if name not in src.sheetnames:
            continue
        for r in rows:
            for c in cols:
                v = src[name].cell(r, c).value
                wb[name].cell(r, c).value = v.strip() if isinstance(v, str) else v

# Тестегі рұқсат етілмеген мәндерді қызылмен белгілеу (оларды қатысушыдан қайта сұрау керек)
from openpyxl.comments import Comment
RED = PatternFill("solid", fgColor="F4B6B6")
bad = 0
for r in range(2, T_LAST + 1):
    for k in range(len(TEST_WORDS)):
        for c, ok in ((3 + 2 * k, (0, 1, 2)), (4 + 2 * k, (0, 1))):
            v = wt.cell(r, c).value
            if v not in (None, "") and v not in ok:
                wt.cell(r, c).fill = RED
                wt.cell(r, c).comment = Comment(f"Бұл бағанға тек {' не '.join(map(str, ok))} жазылады. Қатысушыдан қайта сұраңыз.", "Тексеру")
                bad += 1
if bad:
    print("Тестте рұқсат етілмеген мән:", bad)

if os.path.exists("tuzetuler.json"):  # жетекшімен келісілген түзетулер: {"парақ!ұяшық": мән}
    for ref, val in json.load(open("tuzetuler.json", encoding="utf8")).items():
        sh, cell = ref.split("!")
        wb[sh][cell] = val

for sh in wb.worksheets:
    sh.page_setup.orientation = "landscape"
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
wb.save(OUT)
print("written", OUT)
