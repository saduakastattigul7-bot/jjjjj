"""Толтырылған Zertteu_kestesi.xlsx файлынан нәтижені есептеп, natije.json жазады.
Формулаларға сүйенбейді: бәрі осында қайта есептеледі. БАҚ мониторингі бос болса, natije.json жазылмайды."""
import json, sys
from collections import Counter
from openpyxl import load_workbook
from zertteu import GROUPS, WORDS, STATUS, RARE, TEST_WORDS, GENS, N_PER_GEN

src = sys.argv[1] if len(sys.argv) > 1 else "../Zertteu_kestesi.xlsx"
out = sys.argv[2] if len(sys.argv) > 2 else "natije.json"
wb = load_workbook(src, data_only=True)
s = lambda v: str(v).strip() if v not in (None, "") else ""


def num(v):
    try:
        return int(float(str(v).replace(" ", "").replace(",", ".")))
    except (TypeError, ValueError):
        return None


R = {"n_words": len(WORDS)}
WORDSET = {w for w, *_ in WORDS}

# ---- Сөздер мен БАҚ ----
wsz, wm = wb["Сөздер"], wb["БАҚ мониторингі"]
words, filled = [], 0
for i, (w, g, k, m) in enumerate(WORDS):
    r = i + 2
    h1, c1 = num(wm.cell(r, 3).value), s(wm.cell(r, 4).value).upper()
    if h1 is not None:
        filled += 1
    hits = h1 or 0
    st = "" if h1 is None else "Ұ" if hits < RARE else "Ж" if c1 == "Ж" else "А" if c1 == "А" else "Т"
    words.append({"word": w, "group": g, "kind": k, "meaning": s(wsz.cell(r, 5).value).rstrip(". ") or m,
                  "mark": s(wsz.cell(r, 6).value), "page": s(wsz.cell(r, 7).value),
                  "h1": h1, "c1": c1, "hits": hits, "example": s(wm.cell(r, 5).value),
                  "url": s(wm.cell(r, 6).value), "date": s(wm.cell(r, 7).value)[:10], "status": st})
R["words"] = words
R["media_filled"] = filled
R["marks"] = dict(Counter(x["mark"] for x in words if x["mark"]))
R["by_status"] = {c: sum(x["status"] == c for x in words) for c, _, _ in STATUS}
R["by_group_status"] = {g: {c: sum(x["group"] == g and x["status"] == c for x in words) for c, _, _ in STATUS} for g, _ in GROUPS}
R["by_kind_status"] = {k: {c: sum(x["kind"] == k and x["status"] == c for x in words) for c, _, _ in STATUS} for k in ("Т", "А")}
R["ctx"] = dict(Counter(x["c1"] for x in words if x["c1"]))
R["hits_total"] = sum(x["hits"] for x in words)
R["top_hits"] = [x["word"] for x in sorted(words, key=lambda x: -x["hits"])[:5]]
R["zero_hits"] = [x["word"] for x in words if x["h1"] is not None and x["hits"] == 0]
R["marks_by_kind"] = {k: dict(Counter(x["mark"] for x in words if x["kind"] == k and x["mark"])) for k in ("Т", "А")}

# ---- Тест ----
wt = wb["Тест"]
people = []
for r in range(2, 2 + len(GENS) * N_PER_GEN):
    code = s(wt.cell(r, 1).value)
    vals = [wt.cell(r, 3 + j).value for j in range(2 * len(TEST_WORDS))]
    if all(v in (None, "") for v in vals):
        continue
    people.append({"code": code, "gen": code[:1], "know": [num(vals[2 * k]) for k in range(len(TEST_WORDS))],
                   "use": [num(vals[2 * k + 1]) for k in range(len(TEST_WORDS))]})
R["n_people"] = len(people)
R["n_by_gen"] = {g: sum(p["gen"] == g for p in people) for g, _ in GENS}
pct = lambda a, cond: round(sum(cond(x) for x in a) / len(a) * 100, 1) if a else None
test = {}
for k, w in enumerate(TEST_WORDS):
    test[w] = {}
    for g, _ in GENS:
        kn = [p["know"][k] for p in people if p["gen"] == g and p["know"][k] is not None]
        us = [p["use"][k] for p in people if p["gen"] == g and p["use"][k] is not None]
        test[w][g] = {"ok": pct(kn, lambda x: x == 2), "any": pct(kn, lambda x: x >= 1), "use": pct(us, lambda x: x == 1)}
R["test"] = test
R["gen_avg"] = {}
for g, _ in GENS:
    v = [test[w][g]["ok"] for w in TEST_WORDS if test[w][g]["ok"] is not None]
    R["gen_avg"][g] = round(sum(v) / len(v), 1) if v else None

if not filled:
    print("БАҚ мониторингі бос – natije.json жазылмады.")
else:
    json.dump(R, open(out, "w", encoding="utf8"), ensure_ascii=False, indent=1)
    print("written", out, "words:", filled, "people:", len(people))
