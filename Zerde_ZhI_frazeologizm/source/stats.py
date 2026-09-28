"""Толтырылған Tazhiribe_kestesi.xlsx файлынан нәтижелерді есептеп, natije.json жазады.
Кестеде балл жоқ болса, natije.json жазылмайды (жұмыс мәтінінде [___] орындары қалады)."""
import json, sys
from openpyxl import load_workbook
from tazhiribe import E1, E2, E3, E3_N, AIS, ERR, E3_CODES

src = sys.argv[1] if len(sys.argv) > 1 else "../Tazhiribe_kestesi.xlsx"
out = sys.argv[2] if len(sys.argv) > 2 else "natije.json"
wb = load_workbook(src, data_only=True)
num = lambda v: int(v) if v not in (None, "") and str(v).strip() in "012" else None
txt = lambda v: str(v).strip().upper() if v not in (None, "") else ""

R = {}
ws = wb["1-тәжірибе"]
rows = [[ws.cell(r, c).value for c in range(1, 14)] for r in range(2, len(E1) + 2)]
filled = 0
for k, ai in enumerate(AIS, 1):
    sc = [num(r[5 + 3 * (k - 1)]) for r in rows]
    er = [txt(r[6 + 3 * (k - 1)]) for r in rows]
    filled += sum(s is not None for s in sc)
    for g in ["А", "Ә", "Б", "all"]:
        vals = [s for s, (gg, _, _) in zip(sc, E1) if s is not None and (g == "all" or gg == g)]
        R[f"e1_{g}_{k}"] = round(sum(vals) / (2 * len(vals)) * 100, 1) if vals else None
    R[f"e1_ok_{k}"] = sum(s == 2 for s in sc)
    R[f"e1_half_{k}"] = sum(s == 1 for s in sc)
    R[f"e1_zero_{k}"] = sum(s == 0 for s in sc)
    for code, _, _ in ERR:
        R[f"e1_err_{code}_{k}"] = sum(e == code for e in er)
    R[f"e1_answers_{k}"] = [r[4 + 3 * (k - 1)] for r in rows]
    R[f"e1_scores_{k}"] = sc
for g in ["А", "Ә", "Б", "all"]:
    v = [R[f"e1_{g}_{k}"] for k in range(1, 4) if R[f"e1_{g}_{k}"] is not None]
    R[f"e1_{g}_avg"] = round(sum(v) / len(v), 1) if v else None
for code, _, _ in ERR:
    R[f"e1_err_{code}_sum"] = sum(R[f"e1_err_{code}_{k}"] for k in range(1, 4))

ws = wb["2-тәжірибе"]
rows = [[ws.cell(r, c).value for c in range(1, 14)] for r in range(2, len(E2) + 2)]
for k in range(1, 4):
    sc = [num(r[5 + 3 * (k - 1)]) for r in rows]
    R[f"e2_ant_{k}"] = sum(s or 0 for s in sc[:3])
    R[f"e2_syn_{k}"] = sum(s or 0 for s in sc[3:])
    R[f"e2_all_{k}"] = R[f"e2_ant_{k}"] + R[f"e2_syn_{k}"]
    R[f"e2_answers_{k}"] = [r[4 + 3 * (k - 1)] for r in rows]
    er = [txt(r[6 + 3 * (k - 1)]) for r in rows]
    for code, _, _ in ERR:
        R[f"e2_err_{code}_{k}"] = sum(e == code for e in er)
for code, _, _ in ERR:
    R[f"e2_err_{code}_sum"] = sum(R[f"e2_err_{code}_{k}"] for k in range(1, 4))

ws = wb["3-тәжірибе"]
n3 = len(E3) * E3_N
for k in range(1, 4):
    codes = [txt(ws.cell(r, 2 + 2 * k).value) for r in range(2, n3 + 2)]
    for code, _, _ in E3_CODES:
        R[f"e3_{code}_{k}"] = sum(c == code for c in codes)
    R[f"e3_names_{k}"] = [ws.cell(r, 1 + 2 * k).value for r in range(2, n3 + 2)]
for code, _, _ in E3_CODES:
    R[f"e3_{code}_sum"] = sum(R[f"e3_{code}_{k}"] for k in range(1, 4))

names = wb["Қорытынды"]
R["ai_names"] = [names.cell(2 + k, 2).value or f"ЖИ-{k + 1}" for k in range(3)]

if filled == 0:
    print("Кестеде балл жоқ – natije.json жазылмады.")
else:
    json.dump(R, open(out, "w"), ensure_ascii=False, indent=1)
    print("written", out, "filled scores:", filled)
