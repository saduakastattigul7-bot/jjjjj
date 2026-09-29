"""Оқушы тәжірибе деректерін жазатын Excel кестесін жасайды: Tazhiribe_kestesi.xlsx"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference, Series
from tazhiribe import E1, E2, E3, E3_N, AIS, SCORE, ERR, E3_CODES, PROMPT1, PROMPT2, PROMPT3, GROUPS

OUT = sys.argv[1] if len(sys.argv) > 1 else "Tazhiribe_kestesi.xlsx"
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DCE6F1")
FILL = PatternFill("solid", fgColor="FFF2CC")  # оқушы толтыратын ұяшықтар
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)

wb = Workbook()

# ---------- Нұсқау ----------
ws = wb.active
ws.title = "Нұсқау"
ws.column_dimensions["A"].width = 110
lines = [
    ("ТӘЖІРИБЕ КЕСТЕСІН ТОЛТЫРУ НҰСҚАУЫ", True),
    ("Сары ұяшықтарды ғана толтырыңыз. Қалған ұяшықтар мен «Қорытынды» парағы өздігінен есептеледі.", False),
    ("", False),
    ("1. Үш ЖИ құралын таңдаңыз да, «Қорытынды» парағының жоғарғы жағына атауы мен нұсқасын жазыңыз (ЖИ-1, ЖИ-2, ЖИ-3).", False),
    ("2. Әр сұраныстың алдында жаңа чат ашыңыз (алдыңғы сұрақ жауапқа әсер етпеуі үшін). Сұранысты өзгертпей көшіріп қойыңыз.", False),
    ("3. ЖИ жауабын қысқартып (1–2 сөйлем) «жауабы» бағанына жазыңыз, толық жауаптың скриншотын сақтаңыз (файл атауы: ЖИ1_Т01.png).", False),
    ("4. Жауапты сөздік мағынасымен салыстырып, балл қойыңыз: 2 – дұрыс, 1 – жартылай дұрыс, 0 – қате.", False),
    ("5. Балл 0 немесе 1 болса, қате түрін белгілеңіз: С – сөзбе-сөз, Б – басқа мағына, О – ойдан шығару, Ж – жалпылама.", False),
    ("6. 3-тәжірибеде ЖИ атаған әр тіркесті Кеңесбаев сөздігінен және жұмыстың А қосымшасынан іздеңіз, код қойыңыз: Н, М, Ж, Т.", False),
    ("7. Күмәнді жауапты ғылыми жетекшімен бірге қайта тексеріп, шешімді зерттеу күнделігіне жазыңыз.", False),
    ("8. Сұраныстарға есіміңізді, мектебіңізді және басқа жеке деректерді жазбаңыз.", False),
    ("", False),
    ("Сұраныс үлгілері:", True),
    ("1-тәжірибе: " + PROMPT1.format(x="…"), False),
    ("2-тәжірибе: " + PROMPT2.format(x="…", rel="қарама-қарсы / жақын"), False),
    ("3-тәжірибе: " + PROMPT3.format(x="…"), False),
    ("", False),
    ("Бағалау шкаласы:", True),
    *[(f"{a} – {b}: {c}", False) for a, b, c in SCORE],
    ("Қате түрлері:", True),
    *[(f"{a} – {b}: {c}", False) for a, b, c in ERR],
    ("3-тәжірибе кодтары:", True),
    *[(f"{a} – {b}: {c}", False) for a, b, c in E3_CODES],
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.alignment = WRAP
    c.font = Font(bold=b, size=12 if b else 11)

dv_score = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
dv_err = DataValidation(type="list", formula1='"С,Б,О,Ж"', allow_blank=True)
dv_e3 = DataValidation(type="list", formula1='"Н,М,Ж,Т"', allow_blank=True)


def header(ws, row, titles, widths):
    for j, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row, j, t)
        c.fill, c.font, c.alignment, c.border = HEAD, BOLD, CEN, BOX
        ws.column_dimensions[c.column_letter].width = w


def put(ws, r, j, v=None, fill=False, center=False):
    c = ws.cell(r, j, v)
    c.border = BOX
    c.alignment = CEN if center else WRAP
    if fill:
        c.fill = FILL
    return c


# ---------- 1-тәжірибе ----------
ws = wb.create_sheet("1-тәжірибе")
ws.add_data_validation(dv_score); ws.add_data_validation(dv_err)
t = ["№", "Топ", "Тіркес", "Сөздіктегі мағынасы"]
w = [5, 6, 28, 26]
for a in AIS:
    t += [f"{a} жауабы (қысқаша)", f"{a} балл", f"{a} қате түрі"]
    w += [34, 8, 8]
header(ws, 1, t, w)
ws.freeze_panes = "E2"
for i, (g, x, m) in enumerate(E1, 1):
    r = i + 1
    for j, v in enumerate([i, g, x, m], 1):
        put(ws, r, j, v, center=j <= 2)
    for k in range(len(AIS)):
        base = 5 + 3 * k
        put(ws, r, base, fill=True)
        dv_score.add(put(ws, r, base + 1, fill=True, center=True))
        dv_err.add(put(ws, r, base + 2, fill=True, center=True))
    ws.row_dimensions[r].height = 45

# ---------- 2-тәжірибе ----------
ws = wb.create_sheet("2-тәжірибе")
ws.add_data_validation(dv_score); ws.add_data_validation(dv_err)
t = ["№", "Тапсырма", "Берілген тіркес", "Сөздіктегі жауап"]
w = [5, 11, 26, 30]
for a in AIS:
    t += [f"{a} атаған тіркес", f"{a} балл", f"{a} қате түрі"]
    w += [30, 8, 8]
header(ws, 1, t, w)
for i, (kind, x, ans, m) in enumerate(E2, 1):
    r = i + 1
    for j, v in enumerate([i, kind, x, f"{ans} ({m})"], 1):
        put(ws, r, j, v, center=j <= 2)
    for k in range(len(AIS)):
        base = 5 + 3 * k
        put(ws, r, base, fill=True)
        dv_score.add(put(ws, r, base + 1, fill=True, center=True))
        dv_err.add(put(ws, r, base + 2, fill=True, center=True))
    ws.row_dimensions[r].height = 40
ws.cell(len(E2) + 3, 1, "Балл: 2 – нақты тұрақты тіркес әрі мағынасы дұрыс; 1 – мағынасы дұрыс, бірақ тұрақты тіркес емес (жай сөз); 0 – қате немесе жоқ тіркес.").font = Font(italic=True)

# ---------- 3-тәжірибе ----------
ws = wb.create_sheet("3-тәжірибе")
ws.add_data_validation(dv_e3)
t = ["Мағыналық топ", "№"]
w = [22, 5]
for a in AIS:
    t += [f"{a} атаған тіркес", f"{a} код"]
    w += [32, 8]
header(ws, 1, t, w)
r = 2
for g in E3:
    for n in range(1, E3_N + 1):
        put(ws, r, 1, g if n == 1 else None)
        put(ws, r, 2, n, center=True)
        for k in range(len(AIS)):
            put(ws, r, 3 + 2 * k, fill=True)
            dv_e3.add(put(ws, r, 4 + 2 * k, fill=True, center=True))
        r += 1
    ws.merge_cells(start_row=r - E3_N, start_column=1, end_row=r - 1, end_column=1)
    ws.cell(r - E3_N, 1).alignment = CEN

# ---------- Қорытынды ----------
ws = wb.create_sheet("Қорытынды")
ws.column_dimensions["A"].width = 44
for col in "BCDE":
    ws.column_dimensions[col].width = 16
ws["A1"] = "ЖИ құралдары (атауы, нұсқасы, тексеру күні)"; ws["A1"].font = BOLD
for k, a in enumerate(AIS):
    put(ws, 2 + k, 1, a)
    ws.merge_cells(start_row=2 + k, start_column=2, end_row=2 + k, end_column=5)
    put(ws, 2 + k, 2, fill=True)

cols = ["F", "I", "L"]  # балл бағандары (1- және 2-тәжірибе)
ecols = ["G", "J", "M"]  # қате түрі бағандары
n1 = len(E1)
row = 6
ws.cell(row, 1, "1-тәжірибе: сапа көрсеткіші, % (жинаған балл ÷ ең жоғары балл × 100)").font = BOLD
row += 1
header(ws, row, ["Топ"] + AIS + ["Орташа"], [44, 16, 16, 16, 16])
first_chart_row = row
spans = {"А": (2, 11), "Ә": (12, 21), "Б": (22, 31)}
for g, (a, b) in spans.items():
    row += 1
    put(ws, row, 1, f"{g} – {GROUPS[g].split(' (')[0]}")
    for k, c in enumerate(cols):
        put(ws, row, 2 + k, f"=IFERROR(ROUND(SUM('1-тәжірибе'!{c}{a}:{c}{b})/(2*COUNT('1-тәжірибе'!{c}{a}:{c}{b}))*100,1),\"\")", center=True)
    put(ws, row, 5, f"=IFERROR(ROUND(AVERAGE(B{row}:D{row}),1),\"\")", center=True)
row += 1
put(ws, row, 1, "Барлық 30 тіркес").font = BOLD
for k, c in enumerate(cols):
    put(ws, row, 2 + k, f"=IFERROR(ROUND(SUM('1-тәжірибе'!{c}2:{c}{n1 + 1})/(2*COUNT('1-тәжірибе'!{c}2:{c}{n1 + 1}))*100,1),\"\")", center=True)
put(ws, row, 5, f"=IFERROR(ROUND(AVERAGE(B{row}:D{row}),1),\"\")", center=True)
last_group_row = row - 1

ch = BarChart(); ch.type = "col"; ch.title = "1-тәжірибе: топтар бойынша сапа көрсеткіші, %"
ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=first_chart_row, max_row=last_group_row), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=first_chart_row + 1, max_row=last_group_row))
ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100; ch.height = 8; ch.width = 18
ws.add_chart(ch, "G6")

row += 2
ws.cell(row, 1, "1-тәжірибе: қате түрлерінің саны").font = BOLD
row += 1
header(ws, row, ["Қате түрі"] + AIS + ["Барлығы"], [44, 16, 16, 16, 16])
err_first = row
for code, name, _ in ERR:
    row += 1
    put(ws, row, 1, f"{code} – {name}")
    for k, c in enumerate(ecols):
        put(ws, row, 2 + k, f"=COUNTIF('1-тәжірибе'!{c}2:{c}{n1 + 1},\"{code}\")", center=True)
    put(ws, row, 5, f"=SUM(B{row}:D{row})", center=True)
ch = BarChart(); ch.type = "bar"; ch.title = "Қате түрлері"; ch.grouping = "stacked"; ch.overlap = 100
ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=err_first, max_row=row), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=err_first + 1, max_row=row))
ch.height = 8; ch.width = 18
ws.add_chart(ch, "G24")

row += 2
ws.cell(row, 1, "2-тәжірибе: жинаған балл (ең жоғары – 12)").font = BOLD
row += 1
header(ws, row, ["Тапсырма"] + AIS + ["Орташа"], [44, 16, 16, 16, 16])
for label, a, b in [("Антоним табу (ең жоғары – 6)", 2, 4), ("Синоним табу (ең жоғары – 6)", 5, 7), ("Барлығы (ең жоғары – 12)", 2, 7)]:
    row += 1
    put(ws, row, 1, label)
    for k, c in enumerate(cols):
        put(ws, row, 2 + k, f"=IF(COUNT('2-тәжірибе'!{c}{a}:{c}{b})=0,\"\",SUM('2-тәжірибе'!{c}{a}:{c}{b}))", center=True)
    put(ws, row, 5, f"=IFERROR(ROUND(AVERAGE(B{row}:D{row}),1),\"\")", center=True)

row += 2
ws.cell(row, 1, f"3-тәжірибе: ЖИ атаған {len(E3) * E3_N} тіркестің құрамы").font = BOLD
row += 1
header(ws, row, ["Код"] + AIS + ["Барлығы"], [44, 16, 16, 16, 16])
e3_first = row
last = 1 + len(E3) * E3_N
for code, name, _ in E3_CODES:
    row += 1
    put(ws, row, 1, f"{code} – {name}")
    for k, c in enumerate(["D", "F", "H"]):
        put(ws, row, 2 + k, f"=COUNTIF('3-тәжірибе'!{c}2:{c}{last},\"{code}\")", center=True)
    put(ws, row, 5, f"=SUM(B{row}:D{row})", center=True)
ch = BarChart(); ch.type = "col"; ch.title = "3-тәжірибе: ЖИ атаған тіркестер"; ch.grouping = "percentStacked"; ch.overlap = 100
for rr in range(e3_first + 1, row + 1):
    s = Series(Reference(ws, min_col=2, max_col=4, min_row=rr, max_row=rr), title=ws.cell(rr, 1).value)
    ch.series.append(s)
ch.set_categories(Reference(ws, min_col=2, max_col=4, min_row=e3_first, max_row=e3_first))
ch.height = 8; ch.width = 18
ws.add_chart(ch, "G42")

# Толтырылған ескі кестеден деректерді көшіру: python3 make_xlsx.py жаңа.xlsx толтырылған.xlsx
# (openpyxl сақтағанда диаграммаларды жоғалтады, сондықтан кесте әрдайым осы скриптпен қайта жасалады)
if len(sys.argv) > 2:
    import json, os
    from openpyxl import load_workbook
    src = load_workbook(sys.argv[2], data_only=True)
    one = lambda v: v.strip().upper() if isinstance(v, str) and len(v.strip()) == 1 else (v.strip() if isinstance(v, str) else v)
    for name, rows, cols in [("1-тәжірибе", range(2, 32), range(5, 14)), ("2-тәжірибе", range(2, 8), range(5, 14)),
                             ("3-тәжірибе", range(2, 27), range(3, 9))]:
        for r in rows:
            for c in cols:
                wb[name].cell(r, c).value = one(src[name].cell(r, c).value)
    for r in (2, 3, 4):
        wb["Қорытынды"].cell(r, 2).value = src["Қорытынды"].cell(r, 2).value
    if os.path.exists("tuzetuler.json"):  # жетекшімен келісілген түзетулер: {"парақ!ұяшық": мән}
        for ref, val in json.load(open("tuzetuler.json", encoding="utf8")).items():
            sh, cell = ref.split("!")
            wb[sh][cell] = val

for sh in wb.worksheets:
    sh.page_setup.orientation = "landscape"
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
wb.save(OUT)
print("written", OUT)
