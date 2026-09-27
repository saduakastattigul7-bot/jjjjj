# Зерттеу деректеріне арналған Excel кестесі: формулалар мен диаграммалар өздігінен есептеледі
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
head = Font(bold=True); fill = PatternFill("solid", fgColor="E3EAF3"); inp = PatternFill("solid", fgColor="FFF6CC")
thin = Side(style="thin", color="999999"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
N = 30

def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c); cell.font = head; cell.fill = fill; cell.border = box
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

# 1. Тест
ws = wb.active; ws.title = "Тест"
ws["A1"] = "Ойынға дейінгі және кейінгі тест нәтижелері (балл, 10-нан)"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "Сары ұяшықтарға ойынның «Нәтижелер» бетіндегі деректерді жазыңыз. Аты-жөн емес, код жазылады."
header(ws, 4, ["Оқушы коды", "Ойынға дейін", "Ойыннан кейін", "Өзгеріс"])
dv = DataValidation(type="whole", operator="between", formula1="0", formula2="10", showErrorMessage=True,
                    error="0-ден 10-ға дейінгі сан жазыңыз"); ws.add_data_validation(dv)
for r in range(5, 5 + N):
    ws.cell(row=r, column=1, value=f"3А-{r-4:02d}")
    for c in (2, 3):
        cell = ws.cell(row=r, column=c); cell.fill = inp; dv.add(cell)
    ws.cell(row=r, column=4, value=f'=IF(AND(B{r}<>"",C{r}<>""),C{r}-B{r},"")')
    for c in range(1, 5): ws.cell(row=r, column=c).border = box
e = 5 + N
ws.cell(row=e, column=1, value="Орташа").font = head
for c, L in ((2, "B"), (3, "C"), (4, "D")):
    ws.cell(row=e, column=c, value=f'=IFERROR(ROUND(AVERAGE({L}5:{L}{e-1}),1),"")').font = head
ws.cell(row=e+1, column=1, value="Қатысқандар саны"); ws.cell(row=e+1, column=2, value=f"=COUNT(B5:B{e-1})"); ws.cell(row=e+1, column=3, value=f"=COUNT(C5:C{e-1})")
ws.cell(row=e+2, column=1, value="Нәтижесі өскендер"); ws.cell(row=e+2, column=4, value=f'=COUNTIF(D5:D{e-1},">0")')
ws.cell(row=e+3, column=1, value="Өзгермегендер"); ws.cell(row=e+3, column=4, value=f'=COUNTIF(D5:D{e-1},0)')
ws.cell(row=e+4, column=1, value="Төмендегендер"); ws.cell(row=e+4, column=4, value=f'=COUNTIF(D5:D{e-1},"<0")')
ws["F4"] = "Диаграмма үшін"; ws["F4"].font = head
ws["F5"] = "Ойынға дейін"; ws["G5"] = f"=B{e}"; ws["F6"] = "Ойыннан кейін"; ws["G6"] = f"=C{e}"
ch = BarChart(); ch.title = "Тесттің орташа балы"; ch.y_axis.title = "балл"; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 10
ch.add_data(Reference(ws, min_col=7, min_row=5, max_row=6)); ch.set_categories(Reference(ws, min_col=6, min_row=5, max_row=6))
ch.legend = None; ch.height = 8; ch.width = 12; ws.add_chart(ch, "F9")
for col, w in zip("ABCDEFG", (18, 14, 14, 11, 3, 16, 8)): ws.column_dimensions[col].width = w

# 2. Деңгейлер
ws2 = wb.create_sheet("Деңгейлер")
ws2["A1"] = "Деңгейлер бойынша нәтижелер"; ws2["A1"].font = Font(bold=True, size=13)
ws2["A2"] = "Әр ойыншының әр деңгейдегі қате саны («Нәтижелер» бетіндегі CSV-дан: түрі = деңгей)."
header(ws2, 4, ["Оқушы коды", "1-деңгей", "2-деңгей", "3-деңгей", "4-деңгей", "5-деңгей"])
for r in range(5, 5 + N):
    ws2.cell(row=r, column=1, value=f"3А-{r-4:02d}")
    for c in range(2, 7): ws2.cell(row=r, column=c).fill = inp
    for c in range(1, 7): ws2.cell(row=r, column=c).border = box
ws2.cell(row=e, column=1, value="Орташа қате").font = head
for c in range(2, 7):
    L = "ABCDEF"[c-1]; ws2.cell(row=e, column=c, value=f'=IFERROR(ROUND(AVERAGE({L}5:{L}{e-1}),1),0)').font = head
ch2 = BarChart(); ch2.title = "Деңгейлердегі орташа қате саны"; ch2.y_axis.scaling.min = 0
ch2.add_data(Reference(ws2, min_col=2, max_col=6, min_row=e), from_rows=True, titles_from_data=False)
ch2.set_categories(Reference(ws2, min_col=2, max_col=6, min_row=4)); ch2.legend = None; ch2.height = 8; ch2.width = 14
ws2.add_chart(ch2, "H4"); ws2.column_dimensions["A"].width = 18

# 3. Сауалнама
ws3 = wb.create_sheet("Сауалнама")
ws3["A1"] = "Сауалнама (Иә / Аздап / Жоқ)"; ws3["A1"].font = Font(bold=True, size=13)
qs = ["Ойын саған ұнады ма?", "Ойын арқылы жаңа нәрсе үйрендің бе?", "Қазақ тілі тапсырмаларын осылай тағы орындағың келе ме?"]
header(ws3, 3, ["Оқушы коды"] + [f"{i+1}-сұрақ" for i in range(3)])
dv3 = DataValidation(type="list", formula1='"Иә,Аздап,Жоқ"'); ws3.add_data_validation(dv3)
for r in range(4, 4 + N):
    ws3.cell(row=r, column=1, value=f"3А-{r-3:02d}")
    for c in range(2, 5): cell = ws3.cell(row=r, column=c); cell.fill = inp; dv3.add(cell)
s = 4 + N + 1
header(ws3, s, ["Сұрақ", "Иә", "Аздап", "Жоқ"])
for i, q in enumerate(qs):
    L = "BCD"[i]; ws3.cell(row=s+1+i, column=1, value=q)
    for j, a in enumerate(["Иә", "Аздап", "Жоқ"]):
        ws3.cell(row=s+1+i, column=2+j, value=f'=COUNTIF({L}4:{L}{3+N},"{a}")')
ws3.column_dimensions["A"].width = 52
ch3 = BarChart(); ch3.type = "bar"; ch3.grouping = "stacked"; ch3.overlap = 100; ch3.title = "Сауалнама нәтижелері"
ch3.add_data(Reference(ws3, min_col=2, max_col=4, min_row=s, max_row=s+3), titles_from_data=True)
ch3.set_categories(Reference(ws3, min_col=1, min_row=s+1, max_row=s+3)); ch3.height = 8; ch3.width = 18
ws3.add_chart(ch3, "F3")

wb.save("../12_Derekter_kestesi.xlsx"); print("ok")
