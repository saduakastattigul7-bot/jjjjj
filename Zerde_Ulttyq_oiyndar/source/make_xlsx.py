"""Оқушы толтыратын зерттеу кестесі: Zertteu_kestesi.xlsx
python3 make_xlsx.py шығыс.xlsx [толтырылған_ескі.xlsx]  – ескі кестедегі деректер жаңасына көшіріледі."""
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from zertteu import TYPES, GAMES, CORPUS, FUNCS, KNOW, FOLK, N_EPISODES, N_RESP

OUT = sys.argv[1] if len(sys.argv) > 1 else "Zertteu_kestesi.xlsx"
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DCE6F1")
FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)
TYPE_NAME = dict(TYPES)


def header(ws, row, titles, widths):
    for j, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row, j, t)
        c.fill, c.font, c.alignment, c.border = HEAD, BOLD, CEN, BOX
        if w:
            ws.column_dimensions[c.column_letter].width = w


def put(ws, r, j, v=None, fill=False, center=False):
    c = ws.cell(r, j, v)
    c.border = BOX
    c.alignment = CEN if center else WRAP
    if fill:
        c.fill = FILL
    return c


wb = Workbook()
ws = wb.active
ws.title = "Нұсқау"
ws.column_dimensions["A"].width = 115
lines = [
    ("ЗЕРТТЕУ КЕСТЕСІН ТОЛТЫРУ НҰСҚАУЫ", True),
    ("Сары ұяшықтарды ғана толтырыңыз. «Қорытынды» парағы өздігінен есептеледі.", False),
    ("", False),
    ("1-КЕЗЕҢ. Мәтін талдауы", True),
    ("1. «Шығармалар» парағындағы 15 шығарманы оқыңыз. Қай кітаптан (баспасы, жылы) оқығаныңызды жазып, «Оқылды» бағанына «иә» деп белгілеңіз.", False),
    ("2. Шығармада ұлттық ойын кездессе, «Мәтін талдауы» парағына жаңа жол қосыңыз: бір эпизод – бір жол.", False),
    ("3. Ойынды тізімнен таңдаңыз. Тізімде жоқ болса, «Басқа» деп таңдап, атауын келесі бағанға жазыңыз.", False),
    ("4. Эпизодты 1–2 сөйлеммен өз сөзіңізбен жазыңыз және мәтіннен қысқа үзінді (дәйексөз) мен бетін көрсетіңіз.", False),
    ("5. Ойынның сюжеттегі негізгі қызметін бір кодпен белгілеңіз:", False),
    *[(f"     {a} – {b}: {c}", False) for a, b, c in FUNCS],
    ("6. Ойын тек аталып өтсе де (мысалы, «той болып, бәйге шабылды»), оны жазыңыз – қызметі «Т» болады.", False),
    ("", False),
    ("2-КЕЗЕҢ. Сауалнама", True),
    ("7. Сауалнаманы тек ата-ананың жазбаша келісімі алынған сыныптастар толтырады. Есімдерін жазбаңыз: әр қатысушыға код беріңіз (С01, С02 …).", False),
    ("8. Әр ойын бойынша екі сұрақ: «Бұл ойынды білесің бе?» (0 – білмеймін, 1 – естігенмін, 2 – ойнағанмын/көргенмін) және «Ертегіден не жырдан оқыдың ба?» (1 – иә, 0 – жоқ).", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.alignment = WRAP
    c.font = Font(bold=b, size=12 if b else 11)

# ---- Ойындар (анықтамалық) ----
wg = wb.create_sheet("Ойындар")
header(wg, 1, ["№", "Ойын", "Түрі", "Қысқаша ережесі"], [5, 18, 30, 90])
for i, (g, t, rule) in enumerate(GAMES, 1):
    for j, v in enumerate([i, g, TYPE_NAME[t], rule], 1):
        put(wg, i + 1, j, v, center=j == 1)
put(wg, len(GAMES) + 2, 1, len(GAMES) + 1, center=True)
put(wg, len(GAMES) + 2, 2, "Басқа")
put(wg, len(GAMES) + 2, 3, "Басқа ойындар")
GAME_LAST = len(GAMES) + 2

# ---- Шығармалар ----
wc = wb.create_sheet("Шығармалар")
header(wc, 1, ["№", "Шығарма", "Жанры", "Оқыған басылым (кітап, баспа, жылы)", "Оқылды (иә/жоқ)"], [5, 30, 10, 55, 14])
dv_yes = DataValidation(type="list", formula1='"иә,жоқ"', allow_blank=True)
wc.add_data_validation(dv_yes)
for i, (w, g) in enumerate(CORPUS, 1):
    put(wc, i + 1, 1, i, center=True); put(wc, i + 1, 2, w); put(wc, i + 1, 3, g, center=True)
    put(wc, i + 1, 4, fill=True)
    dv_yes.add(put(wc, i + 1, 5, fill=True, center=True))
CORP_LAST = len(CORPUS) + 1

# ---- Мәтін талдауы ----
we = wb.create_sheet("Мәтін талдауы")
header(we, 1, ["№", "Шығарма", "Жанры", "Ойын", "Басқа ойын атауы", "Ойын түрі", "Кейіпкер", "Эпизод (өз сөзіңізбен)", "Қызметі (код)", "Дәйексөз (қысқа үзінді)", "Беті"],
       [5, 26, 9, 16, 16, 24, 16, 48, 10, 40, 7])
we.freeze_panes = "B2"
dv_work = DataValidation(type="list", formula1=f"=Шығармалар!$B$2:$B${CORP_LAST}", allow_blank=True)
dv_game = DataValidation(type="list", formula1=f"=Ойындар!$B$2:$B${GAME_LAST}", allow_blank=True)
dv_func = DataValidation(type="list", formula1='"' + ",".join(f[0] for f in FUNCS) + '"', allow_blank=True)
for dv in (dv_work, dv_game, dv_func):
    we.add_data_validation(dv)
for i in range(1, N_EPISODES + 1):
    r = i + 1
    put(we, r, 1, i, center=True)
    dv_work.add(put(we, r, 2, fill=True))
    put(we, r, 3, f'=IFERROR(VLOOKUP(B{r},Шығармалар!$B$2:$C${CORP_LAST},2,0),"")', center=True)
    dv_game.add(put(we, r, 4, fill=True))
    put(we, r, 5, fill=True)
    put(we, r, 6, f'=IFERROR(VLOOKUP(D{r},Ойындар!$B$2:$C${GAME_LAST},2,0),"")')
    put(we, r, 7, fill=True); put(we, r, 8, fill=True)
    dv_func.add(put(we, r, 9, fill=True, center=True))
    put(we, r, 10, fill=True); put(we, r, 11, fill=True, center=True)
EP_LAST = N_EPISODES + 1

# ---- Сауалнама ----
wsv = wb.create_sheet("Сауалнама")
titles = ["Код", "Сынып"]
for g, _, _ in GAMES:
    titles += [f"{g}: білу (0/1/2)", f"{g}: ертегіден (1/0)"]
header(wsv, 1, titles, [7, 7] + [11] * (2 * len(GAMES)))
wsv.row_dimensions[1].height = 60
wsv.freeze_panes = "C2"
dv_k = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
dv_f = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
wsv.add_data_validation(dv_k); wsv.add_data_validation(dv_f)
for i in range(1, N_RESP + 1):
    r = i + 1
    put(wsv, r, 1, f"С{i:02d}", center=True)
    put(wsv, r, 2, fill=True, center=True)
    for k in range(len(GAMES)):
        dv_k.add(put(wsv, r, 3 + 2 * k, fill=True, center=True))
        dv_f.add(put(wsv, r, 4 + 2 * k, fill=True, center=True))
SV_LAST = N_RESP + 1

# ---- Қорытынды ----
wq = wb.create_sheet("Қорытынды")
for col, w in zip("ABCDEF", [34, 14, 14, 14, 14, 14]):
    wq.column_dimensions[col].width = w
row = 1
wq.cell(row, 1, "Мәтін талдауы: ойындардың кездесуі").font = BOLD
row += 1
header(wq, row, ["Ойын", "Ертегіде", "Жырда", "Барлығы"], [None] * 4)
g_first = row
for g, _, _ in GAMES + [("Басқа", "", "")]:
    row += 1
    put(wq, row, 1, g)
    put(wq, row, 2, f"=COUNTIFS('Мәтін талдауы'!$D$2:$D${EP_LAST},A{row},'Мәтін талдауы'!$C$2:$C${EP_LAST},\"Ертегі\")", center=True)
    put(wq, row, 3, f"=COUNTIFS('Мәтін талдауы'!$D$2:$D${EP_LAST},A{row},'Мәтін талдауы'!$C$2:$C${EP_LAST},\"Жыр\")", center=True)
    put(wq, row, 4, f"=B{row}+C{row}", center=True)
g_last = row
row += 1
put(wq, row, 1, "Барлық эпизод").font = BOLD
for j, col in ((2, "B"), (3, "C"), (4, "D")):
    put(wq, row, j, f"=SUM({col}{g_first + 1}:{col}{g_last})", center=True)
ch = BarChart(); ch.type = "bar"; ch.title = "Ойындардың шығармаларда кездесуі"; ch.grouping = "stacked"; ch.overlap = 100
ch.add_data(Reference(wq, min_col=2, max_col=3, min_row=g_first, max_row=g_last), titles_from_data=True)
ch.set_categories(Reference(wq, min_col=1, min_row=g_first + 1, max_row=g_last)); ch.height = 9; ch.width = 16
wq.add_chart(ch, "H1")

row += 2
wq.cell(row, 1, "Ойын түрлері").font = BOLD
row += 1
header(wq, row, ["Түрі", "Ертегіде", "Жырда", "Барлығы"], [None] * 4)
for _, tname in TYPES:
    row += 1
    put(wq, row, 1, tname)
    put(wq, row, 2, f"=COUNTIFS('Мәтін талдауы'!$F$2:$F${EP_LAST},A{row},'Мәтін талдауы'!$C$2:$C${EP_LAST},\"Ертегі\")", center=True)
    put(wq, row, 3, f"=COUNTIFS('Мәтін талдауы'!$F$2:$F${EP_LAST},A{row},'Мәтін талдауы'!$C$2:$C${EP_LAST},\"Жыр\")", center=True)
    put(wq, row, 4, f"=B{row}+C{row}", center=True)

row += 2
wq.cell(row, 1, "Ойынның сюжеттегі қызметі").font = BOLD
row += 1
header(wq, row, ["Қызметі", "Ертегіде", "Жырда", "Барлығы"], [None] * 4)
f_first = row
for code, name, _ in FUNCS:
    row += 1
    put(wq, row, 1, f"{code} – {name}")
    put(wq, row, 2, f"=COUNTIFS('Мәтін талдауы'!$I$2:$I${EP_LAST},\"{code}\",'Мәтін талдауы'!$C$2:$C${EP_LAST},\"Ертегі\")", center=True)
    put(wq, row, 3, f"=COUNTIFS('Мәтін талдауы'!$I$2:$I${EP_LAST},\"{code}\",'Мәтін талдауы'!$C$2:$C${EP_LAST},\"Жыр\")", center=True)
    put(wq, row, 4, f"=B{row}+C{row}", center=True)
ch = BarChart(); ch.type = "col"; ch.title = "Ойынның сюжеттегі қызметі"
ch.add_data(Reference(wq, min_col=2, max_col=3, min_row=f_first, max_row=row), titles_from_data=True)
ch.set_categories(Reference(wq, min_col=1, min_row=f_first + 1, max_row=row)); ch.height = 8; ch.width = 16
wq.add_chart(ch, "H20")

row += 2
wq.cell(row, 1, "Сауалнама: қатысушылар саны").font = BOLD
put(wq, row, 2, f"=COUNTA(Сауалнама!$B$2:$B${SV_LAST})", center=True)
row += 1
header(wq, row, ["Ойын", "Біледі, %", "Ойнаған, %", "Ертегіден оқыған, %"], [None] * 4)
s_first = row
from openpyxl.utils import get_column_letter
for k, (g, _, _) in enumerate(GAMES):
    row += 1
    ck, cf = get_column_letter(3 + 2 * k), get_column_letter(4 + 2 * k)
    rng_k, rng_f = f"Сауалнама!${ck}$2:${ck}${SV_LAST}", f"Сауалнама!${cf}$2:${cf}${SV_LAST}"
    put(wq, row, 1, g)
    put(wq, row, 2, f'=IFERROR(ROUND(COUNTIF({rng_k},">=1")/COUNT({rng_k})*100,1),"")', center=True)
    put(wq, row, 3, f'=IFERROR(ROUND(COUNTIF({rng_k},2)/COUNT({rng_k})*100,1),"")', center=True)
    put(wq, row, 4, f'=IFERROR(ROUND(COUNTIF({rng_f},1)/COUNT({rng_f})*100,1),"")', center=True)
ch = BarChart(); ch.type = "bar"; ch.title = "Сауалнама: ойындарды білу, %"
ch.add_data(Reference(wq, min_col=2, max_col=4, min_row=s_first, max_row=row), titles_from_data=True)
ch.set_categories(Reference(wq, min_col=1, min_row=s_first + 1, max_row=row)); ch.height = 11; ch.width = 16
wq.add_chart(ch, "H37")

# ---- Ескі кестеден көшіру ----
if len(sys.argv) > 2:
    src = load_workbook(sys.argv[2], data_only=True)
    for name, rows, cols in [("Шығармалар", range(2, CORP_LAST + 1), (4, 5)),
                             ("Мәтін талдауы", range(2, EP_LAST + 1), (2, 4, 5, 7, 8, 9, 10, 11)),
                             ("Сауалнама", range(2, SV_LAST + 1), range(2, 3 + 2 * len(GAMES)))]:
        for r in rows:
            for c in cols:
                v = src[name].cell(r, c).value
                wb[name].cell(r, c).value = v.strip() if isinstance(v, str) else v

for sh in wb.worksheets:
    sh.page_setup.orientation = "landscape"
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
wb.save(OUT)
print("written", OUT)
