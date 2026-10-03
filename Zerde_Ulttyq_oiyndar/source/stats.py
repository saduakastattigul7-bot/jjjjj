"""Толтырылған Zertteu_kestesi.xlsx файлынан нәтижені есептеп, natije.json жазады.
Мәтін талдауында эпизод жоқ болса, natije.json жазылмайды."""
import json, sys
from collections import Counter
from openpyxl import load_workbook
from zertteu import TYPES, GAMES, CORPUS, FUNCS, N_EPISODES, N_RESP

src = sys.argv[1] if len(sys.argv) > 1 else "../Zertteu_kestesi.xlsx"
out = sys.argv[2] if len(sys.argv) > 2 else "natije.json"
wb = load_workbook(src, data_only=True)
s = lambda v: str(v).strip() if v not in (None, "") else ""
GENRE = dict(CORPUS)
GTYPE = {g: t for g, t, _ in GAMES}
TYPE_NAME = dict(TYPES)
FCODES = [f[0] for f in FUNCS]

R = {}
wc = wb["Шығармалар"]
R["works"] = [{"name": wc.cell(r, 2).value, "genre": wc.cell(r, 3).value, "edition": s(wc.cell(r, 4).value),
               "read": s(wc.cell(r, 5).value).lower() == "иә"} for r in range(2, len(CORPUS) + 2)]
R["works_read"] = sum(w["read"] for w in R["works"])

we = wb["Мәтін талдауы"]
eps = []
for r in range(2, N_EPISODES + 2):
    work, game = s(we.cell(r, 2).value), s(we.cell(r, 4).value)
    if not work or not game:
        continue
    other = s(we.cell(r, 5).value)
    eps.append({
        "work": work, "genre": GENRE.get(work, ""), "game": game if game != "Басқа" else (other or "Басқа"),
        "listed": game != "Басқа", "type": GTYPE.get(game, "БАСҚА"), "hero": s(we.cell(r, 7).value),
        "episode": s(we.cell(r, 8).value), "func": s(we.cell(r, 9).value).upper(), "quote": s(we.cell(r, 10).value),
        "page": s(we.cell(r, 11).value),
    })
R["episodes"] = eps
R["ep_total"] = len(eps)
for gname in ("Ертегі", "Жыр"):
    R[f"ep_{gname}"] = sum(e["genre"] == gname for e in eps)
R["by_game"] = {g: {"Ертегі": sum(e["game"] == g and e["genre"] == "Ертегі" for e in eps),
                    "Жыр": sum(e["game"] == g and e["genre"] == "Жыр" for e in eps)} for g, _, _ in GAMES}
R["other_games"] = sorted({e["game"] for e in eps if not e["listed"]})
R["by_type"] = {t: {"Ертегі": sum(e["type"] == t and e["genre"] == "Ертегі" for e in eps),
                    "Жыр": sum(e["type"] == t and e["genre"] == "Жыр" for e in eps)} for t, _ in TYPES}
R["by_func"] = {f: {"Ертегі": sum(e["func"] == f and e["genre"] == "Ертегі" for e in eps),
                    "Жыр": sum(e["func"] == f and e["genre"] == "Жыр" for e in eps)} for f in FCODES}
R["works_with_games"] = len({e["work"] for e in eps})
R["games_found"] = len({e["game"] for e in eps})
R["by_work"] = dict(Counter(e["work"] for e in eps))

wsv = wb["Сауалнама"]
resp = []
for r in range(2, N_RESP + 2):
    vals = [wsv.cell(r, 3 + j).value for j in range(2 * len(GAMES))]
    if all(v in (None, "") for v in vals):
        continue
    resp.append({"class": s(wsv.cell(r, 2).value), "know": [vals[2 * k] for k in range(len(GAMES))],
                 "folk": [vals[2 * k + 1] for k in range(len(GAMES))]})
R["n_resp"] = len(resp)
num = lambda v: int(v) if str(v).strip() in ("0", "1", "2") else None
survey = {}
for k, (g, _, _) in enumerate(GAMES):
    kn = [num(p["know"][k]) for p in resp if num(p["know"][k]) is not None]
    fo = [num(p["folk"][k]) for p in resp if num(p["folk"][k]) is not None]
    survey[g] = {
        "know": round(sum(x >= 1 for x in kn) / len(kn) * 100, 1) if kn else None,
        "played": round(sum(x == 2 for x in kn) / len(kn) * 100, 1) if kn else None,
        "folk": round(sum(x == 1 for x in fo) / len(fo) * 100, 1) if fo else None,
    }
R["survey"] = survey

if not eps:
    print("Мәтін талдауында эпизод жоқ – natije.json жазылмады.")
else:
    json.dump(R, open(out, "w", encoding="utf8"), ensure_ascii=False, indent=1)
    print("written", out, "episodes:", len(eps), "respondents:", len(resp))
