"""Data workbook: the student types raw counts/scores, percentages and charts update automatically."""
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

out = sys.argv[1] if len(sys.argv) > 1 else "08_Derekter.xlsx"
wb = Workbook()
B = Font(bold=True)
IN = PatternFill("solid", fgColor="FFF2A8")    # cells the student fills in
HD = PatternFill("solid", fgColor="E7EEF5")
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def head(ws, row, values):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font, cell.fill, cell.border = B, HD, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


# ---- Sheet 1: survey ----
ws = wb.active
ws.title = "Сауалнама"
ws["A1"] = "Сауалнама нәтижелері (жұмыстағы 5-кесте)"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "Сары ұяшықтарға ғана сан жаз. Пайыздар өзі есептеледі."
ws["A3"] = "Барлық қатысушы саны:"; ws["A3"].font = B
ws["C3"].fill = IN; ws["C3"].border = BOX
head(ws, 5, ["Сұрақ", "Жауап", "Саны", "%"])
Q = [
    ("1. Қазақ ертегілерін оқығанды ұнатасың ба?", ["Иә", "Кейде", "Жоқ"]),
    ("2. Соңғы бір айда кітаптан ертегі оқыдың ба?", ["Иә", "Жоқ"]),
    ("3. «Ер Төстік» ертегісін қалай білесің?", ["Кітаптан оқыдым", "Мультфильмін көрдім", "Атын ғана естідім", "Білмеймін"]),
    ("4. Кітапты оқуға сені не қызықтырады? (бірнешеуі)", ["Мұқабасы мен суреті", "Досымның кеңесі", "Мультфильм немесе бейне", "Мұғалімнің тапсырмасы"]),
    ("5. Жарнаманы көбіне қай жерден көресің?", ["Телефоннан", "Теледидардан", "Көшедегі плакаттан", "Мектептегі хабарландырудан"]),
]
r = 6
q3 = None
for q, answers in Q:
    if q.startswith("3."):
        q3 = (r, r + len(answers) - 1)
    for i, a in enumerate(answers):
        ws.cell(row=r, column=1, value=q if i == 0 else None)
        ws.cell(row=r, column=2, value=a)
        c = ws.cell(row=r, column=3); c.fill = IN
        ws.cell(row=r, column=4, value=f'=IF(OR($C$3="",C{r}=""),"",ROUND(C{r}/$C$3*100,0))')
        for col in range(1, 5):
            ws.cell(row=r, column=col).border = BOX
        r += 1
ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 10; ws.column_dimensions["D"].width = 10

# ---- Sheet 2: slogans and posters ----
s2 = wb.create_sheet("Слоган және плакат")
s2["A1"] = "Слогандар (4-кесте) және плакаттарды салыстыру (6-кесте)"; s2["A1"].font = Font(bold=True, size=13)
head(s2, 3, ["Слоган", "Таңдағандар саны", "%"])
slogans = ["«Жер астына түсіп, жеңіп шыққан батыр – Ер Төстік!»", "«Шалқұйрық шапса – ертегі басталады!»",
           "«Бір кітап – мың шытырман оқиға!»", "«Ер Төстік – оқы да, батыр бол!»"]
for i, t in enumerate(slogans):
    rr = 4 + i
    s2.cell(row=rr, column=1, value=t); s2.cell(row=rr, column=2).fill = IN
    s2.cell(row=rr, column=3, value=f'=IF(OR(B{rr}="",$B$8=0),"",ROUND(B{rr}/$B$8*100,0))')
s2["A8"] = "Барлығы"; s2["A8"].font = B; s2["B8"] = "=SUM(B4:B7)"
head(s2, 11, ["Көрсеткіш", "А плакаты", "Б плакаты"])
rows = ["«Оқығым келеді» деп таңдағандар саны", "Сол, %", "10 минуттан кейін кейіпкердің атын есте сақтағандар",
        "10 минуттан кейін ертегінің атын есте сақтағандар"]
for i, t in enumerate(rows):
    rr = 12 + i
    s2.cell(row=rr, column=1, value=t)
    if t == "Сол, %":
        s2.cell(row=rr, column=2, value='=IF(SUM(B12:C12)=0,"",ROUND(B12/SUM(B12:C12)*100,0))')
        s2.cell(row=rr, column=3, value='=IF(SUM(B12:C12)=0,"",ROUND(C12/SUM(B12:C12)*100,0))')
    else:
        s2.cell(row=rr, column=2).fill = IN; s2.cell(row=rr, column=3).fill = IN
for row in s2.iter_rows(min_row=3, max_row=15, max_col=3):
    for c in row:
        if c.row not in (9, 10):
            c.border = BOX
s2.column_dimensions["A"].width = 55; s2.column_dimensions["B"].width = 18; s2.column_dimensions["C"].width = 18

# ---- Sheet 3: before/after scores ----
s3 = wb.create_sheet("Дейін-кейін")
s3["A1"] = "«Ер Төстікті оқығың келе ме?» (1–5 балл) – 7-кесте"; s3["A1"].font = Font(bold=True, size=13)
s3["A2"] = "Әр оқушының балын өз жолына жаз (аты-жөні жазылмайды, тек нөмір). Бір аптадан кейін оқыса – 1, оқымаса – 0."
head(s3, 4, ["Парақ №", "Дейін (1–5)", "Кейін (1–5)", "Бір аптада оқыды (1/0)"])
N = 35
for i in range(N):
    rr = 5 + i
    s3.cell(row=rr, column=1, value=i + 1)
    for col in (2, 3, 4):
        s3.cell(row=rr, column=col).fill = IN
last = 4 + N
summary = [
    ("Қатысушылар саны", f"=COUNT(B5:B{last})", f"=COUNT(C5:C{last})"),
    ("Орташа балл", f'=IF(COUNT(B5:B{last})=0,"",ROUND(AVERAGE(B5:B{last}),2))', f'=IF(COUNT(C5:C{last})=0,"",ROUND(AVERAGE(C5:C{last}),2))'),
    ("«5» қойғандар саны", f"=COUNTIF(B5:B{last},5)", f"=COUNTIF(C5:C{last},5)"),
    ("«4» немесе «5» қойғандар, %", f'=IF(COUNT(B5:B{last})=0,"",ROUND((COUNTIF(B5:B{last},4)+COUNTIF(B5:B{last},5))/COUNT(B5:B{last})*100,0))',
     f'=IF(COUNT(C5:C{last})=0,"",ROUND((COUNTIF(C5:C{last},4)+COUNTIF(C5:C{last},5))/COUNT(C5:C{last})*100,0))'),
    ("Бір аптада оқығандар саны", "", f"=SUM(D5:D{last})"),
]
head(s3, 4, ["Парақ №", "Дейін (1–5)", "Кейін (1–5)", "Бір аптада оқыды (1/0)"])
s3["F4"], s3["G4"], s3["H4"] = "Көрсеткіш", "Дейін", "Кейін"
for c in ("F4", "G4", "H4"):
    s3[c].font, s3[c].fill = B, HD
for i, (t, a, b) in enumerate(summary):
    s3.cell(row=5 + i, column=6, value=t)
    s3.cell(row=5 + i, column=7, value=a or None)
    s3.cell(row=5 + i, column=8, value=b)
# distribution of scores for chart 2
s3["F12"], s3["G12"], s3["H12"] = "Балл", "Дейін (оқушы саны)", "Кейін (оқушы саны)"
for c in ("F12", "G12", "H12"):
    s3[c].font, s3[c].fill = B, HD
for k in range(1, 6):
    rr = 12 + k
    s3.cell(row=rr, column=6, value=str(k))
    s3.cell(row=rr, column=7, value=f"=COUNTIF(B5:B{last},{k})")
    s3.cell(row=rr, column=8, value=f"=COUNTIF(C5:C{last},{k})")
for col, w in zip("ABCDEFGH", (9, 12, 12, 20, 3, 32, 18, 18)):
    s3.column_dimensions[col].width = w

# ---- Sheet 4: charts ----
s4 = wb.create_sheet("Диаграммалар")
s4["A1"] = "Диаграммалар өзі жаңарады. Тінтуірдің оң жағымен басып, «Көшіру» → Word құжатына қой."
c1 = BarChart(); c1.type = "col"; c1.style = 10
c1.title = "1-диаграмма. «Ер Төстік» ертегісін қалай білесің? (%)"
c1.add_data(Reference(ws, min_col=4, min_row=q3[0], max_row=q3[1]), titles_from_data=False)
c1.set_categories(Reference(ws, min_col=2, min_row=q3[0], max_row=q3[1]))
c1.legend = None; c1.y_axis.title = "%"; c1.height, c1.width = 9, 17
s4.add_chart(c1, "A3")
c2 = BarChart(); c2.type = "col"; c2.style = 10
c2.title = "2-диаграмма. «Ер Төстікті оқығың келе ме?» (оқушы саны)"
c2.add_data(Reference(s3, min_col=7, max_col=8, min_row=12, max_row=17), titles_from_data=True)
c2.set_categories(Reference(s3, min_col=6, min_row=13, max_row=17))
c2.x_axis.title = "Балл (1 – мүлдем келмейді, 5 – өте қатты келеді)"; c2.y_axis.title = "Оқушы саны"
c2.height, c2.width = 9, 17
s4.add_chart(c2, "A23")
wb.save(out)
print("written", out)
